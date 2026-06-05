
from __future__ import annotations

import json
import re
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from .artifact_acceptance_models import (
    ArtifactAcceptanceReport,
    ArtifactAcceptanceRequest,
    ArtifactFinding,
    advisory_artifact_findings,
    artifact_ref_payload,
    kind_for_path,
)
from .artifact_binary_signature import binary_signature_finding
from .artifact_capabilities import artifact_capability
from .artifact_collection_contract import collection_contract_findings
from .artifact_csv_acceptance import validate_csv_artifact
from .artifact_document_acceptance import (
    document_quality_artifact_findings,
    validate_docx_artifact,
    validate_text_artifact,
)
from .artifact_html_contract import html_contract_findings, record_resource_ref
from .artifact_html_refs import image_ref_findings, scan_html_refs
from .artifact_staged_evidence import staged_source_evidence_findings
from .artifact_static_site_contract import validate_static_site_artifact
from .artifact_structured_contracts import (
    json_contract_findings,
    markdown_section_findings,
    text_size_findings,
)
from .artifact_validator_registry import (
    ArtifactValidator,
    resolve_artifact_validator,
)
from .artifact_xlsx_contract import xlsx_contract_findings


def validate_html_artifact(request: ArtifactAcceptanceRequest) -> ArtifactAcceptanceReport:
    path = Path(request.path)
    if _outside_workspace(path, request.workspace_root):
        return _outside_workspace_report(path)
    if not path.exists():
        finding = ArtifactFinding(
            code="ARTIFACT_MISSING",
            severity="hard",
            message="HTML artifact does not exist.",
            location=str(path),
        )
        return ArtifactAcceptanceReport(ok=False, artifact_ref=str(path), findings=[finding])
    text = path.read_text(encoding="utf-8", errors="replace")
    refs = scan_html_refs(text)
    findings = _finding_records(
        [
            *image_ref_findings(refs, path=path, workspace_root=request.workspace_root),
            *html_contract_findings(text, refs.resources, request.validation_contract),
        ]
    )
    return ArtifactAcceptanceReport(
        ok=not any(item.severity == "hard" for item in findings),
        artifact_ref=str(path),
        artifact_kind="html",
        findings=findings,
    )


def validate_artifact(request: ArtifactAcceptanceRequest) -> ArtifactAcceptanceReport:
    path = Path(request.path)
    if _outside_workspace(path, request.workspace_root):
        return _outside_workspace_report(path)
    if not path.exists():
        return _missing_report(path, kind=_request_capability(request).kind)
    validator = resolve_artifact_validator(
        request,
        named_validators=_named_validators(),
        kind_validators=_kind_validators(),
        default_validator=_validate_generic_request,
    )
    return validator(request)


def _named_validators() -> dict[str, ArtifactValidator]:
    return {
        "artifact_acceptance": validate_by_artifact_kind,
        "static_site_check": validate_static_site_artifact,
        "spreadsheet_acceptance": _validate_xlsx_request,
        "document_acceptance": _validate_pdf_request,
    }


def _kind_validators() -> dict[str, ArtifactValidator]:
    return {
        "html": validate_html_artifact,
        "htm": validate_html_artifact,
        "json": _validate_json_request,
        "md": _validate_markdown_request,
        "markdown": _validate_markdown_request,
        "txt": _validate_text_request,
        "csv": _validate_csv_request,
        "xlsx": _validate_xlsx_request,
        "pdf": _validate_pdf_request,
        "docx": _validate_docx_request,
    }


def validate_by_artifact_kind(request: ArtifactAcceptanceRequest) -> ArtifactAcceptanceReport:
    capability = _request_capability(request)
    return _kind_validators().get(capability.validator_key or capability.kind, _validate_generic_request)(request)


def _request_capability(request: ArtifactAcceptanceRequest):
    validation = request.validation_contract or {}
    declared_kind = str(validation.get("artifact_kind") or validation.get("kind") or "").strip()
    declared_mime = str(validation.get("mime_type") or validation.get("mime") or "").strip()
    return artifact_capability(request.path, declared_kind=declared_kind, declared_mime=declared_mime)


def _missing_report(path: Path, *, kind: str) -> ArtifactAcceptanceReport:
    finding = ArtifactFinding(
        code="ARTIFACT_MISSING",
        severity="hard",
        message="Artifact does not exist.",
        location=str(path),
    )
    return ArtifactAcceptanceReport(ok=False, artifact_ref=str(path), artifact_kind=kind, findings=[finding])


def _outside_workspace(path: Path, workspace_root: Path | None) -> bool:
    if workspace_root is None:
        return False
    try:
        path.resolve(strict=False).relative_to(Path(workspace_root).resolve(strict=False))
        return False
    except ValueError:
        return True


def validation_workspace_root_for_item(
    item: dict[str, object],
    path: Path,
    workspace_root: Path,
) -> Path:
    resolved_path = path.expanduser().resolve(strict=False)
    workspace = workspace_root.expanduser().resolve(strict=False)
    if _path_is_under(resolved_path, workspace):
        return workspace
    return _matching_allowed_output_root(item, resolved_path, workspace) or workspace


def _matching_allowed_output_root(item: dict[str, object], resolved_path: Path, workspace: Path) -> Path | None:
    roots = item.get("allowed_output_roots")
    if not isinstance(roots, list):
        return None
    for raw in roots:
        root = _resolve_allowed_output_root(raw, workspace, resolved_path)
        if root is not None and _path_is_under(resolved_path, root):
            return root
    return None


def _resolve_allowed_output_root(raw: object, workspace: Path, resolved_path: Path) -> Path | None:
    text = str(raw or "").strip()
    if not text:
        return None
    try:
        candidate = Path(text).expanduser()
        root = candidate.resolve(strict=False) if candidate.is_absolute() else (workspace / candidate).resolve(strict=False)
    except (OSError, RuntimeError):
        return None
    if root == resolved_path and root.suffix:
        return root.parent
    return root


def _path_is_under(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def _outside_workspace_report(path: Path) -> ArtifactAcceptanceReport:
    finding = ArtifactFinding("ARTIFACT_PATH_OUTSIDE_WORKSPACE", "hard", "Artifact path is outside workspace_root.", str(path))
    return ArtifactAcceptanceReport(ok=False, artifact_ref=str(path), artifact_kind=kind_for_path(path), findings=[finding])


def _validate_json_request(request: ArtifactAcceptanceRequest) -> ArtifactAcceptanceReport:
    return _validate_json(Path(request.path), request.validation_contract)


def _validate_markdown_request(request: ArtifactAcceptanceRequest) -> ArtifactAcceptanceReport:
    return _validate_markdown(Path(request.path), request.validation_contract)


def _validate_text_request(request: ArtifactAcceptanceRequest) -> ArtifactAcceptanceReport:
    return validate_text_artifact(Path(request.path), request.validation_contract, workspace_root=request.workspace_root)


def _validate_csv_request(request: ArtifactAcceptanceRequest) -> ArtifactAcceptanceReport:
    return validate_csv_artifact(Path(request.path), request.validation_contract)


def _validate_xlsx_request(request: ArtifactAcceptanceRequest) -> ArtifactAcceptanceReport:
    return _validate_xlsx(Path(request.path), request.validation_contract, workspace_root=request.workspace_root)


def _validate_pdf_request(request: ArtifactAcceptanceRequest) -> ArtifactAcceptanceReport:
    return _validate_pdf(
        Path(request.path),
        request.validation_contract,
        workspace_root=request.workspace_root,
    )


def _validate_docx_request(request: ArtifactAcceptanceRequest) -> ArtifactAcceptanceReport:
    return validate_docx_artifact(Path(request.path), request.validation_contract, workspace_root=request.workspace_root)


def _validate_generic_request(request: ArtifactAcceptanceRequest) -> ArtifactAcceptanceReport:
    return _validate_generic(Path(request.path))


def _validate_json(path: Path, validation_contract: dict[str, object] | None = None) -> ArtifactAcceptanceReport:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        finding = ArtifactFinding(
            code="JSON_INVALID",
            severity="hard",
            message=f"Invalid JSON: {exc.msg}",
            location=str(exc.pos),
        )
        return _report_with_finding(path, "json", finding)
    if not isinstance(value, (dict, list)):
        finding = ArtifactFinding(
            code="JSON_UNEXPECTED_TOP_LEVEL",
            severity="hard",
            message="JSON top-level must be object or array.",
        )
        return _report_with_finding(path, "json", finding)
    findings = json_contract_findings(path, value, validation_contract or {})
    return ArtifactAcceptanceReport(
        ok=not any(item.severity == "hard" for item in findings),
        artifact_ref=str(path),
        artifact_kind="json",
        findings=findings,
    )


def _validate_markdown(
    path: Path,
    validation_contract: dict[str, object] | None = None,
) -> ArtifactAcceptanceReport:
    text = path.read_text(encoding="utf-8", errors="replace")
    if not text:
        finding = ArtifactFinding(code="ARTIFACT_EMPTY", severity="hard", message="Artifact is empty.")
        return _report_with_finding(path, "md", finding)
    findings = advisory_artifact_findings(
        [
            *text_size_findings(path, text, validation_contract or {}),
            *markdown_section_findings(path, text, validation_contract or {}),
        ]
    )
    findings.extend(_required_text_findings(path, text, validation_contract or {}))
    findings.extend(_forbidden_text_findings(path, text, validation_contract or {}))
    findings.extend(document_quality_artifact_findings(path, validation_contract, workspace_root=path.parent))
    return ArtifactAcceptanceReport(
        ok=not any(item.severity == "hard" for item in findings),
        artifact_ref=str(path),
        artifact_kind="md",
        findings=findings,
    )


def _required_text_findings(path: Path, text: str, validation_contract: dict[str, object]) -> list[ArtifactFinding]:
    findings: list[ArtifactFinding] = []
    for token in _required_string_items(validation_contract):
        if token not in text:
            findings.append(
                ArtifactFinding(
                    code="ARTIFACT_REQUIRED_TEXT_MISSING",
                    severity="hard",
                    message="Artifact text is missing a required string from validation_contract.",
                    location=str(path),
                    value=token[:200],
                )
            )
    for pattern in _required_regex_items(validation_contract):
        try:
            matched = re.search(pattern, text, flags=re.MULTILINE) is not None
        except re.error as exc:
            findings.append(
                ArtifactFinding(
                    code="ARTIFACT_REQUIRED_REGEX_INVALID",
                    severity="hard",
                    message=f"validation_contract required regex is invalid: {exc}",
                    location=str(path),
                    value=pattern[:200],
                )
            )
            continue
        if not matched:
            findings.append(
                ArtifactFinding(
                    code="ARTIFACT_REQUIRED_REGEX_MISSING",
                    severity="hard",
                    message="Artifact text does not match a required regex from validation_contract.",
                    location=str(path),
                    value=pattern[:200],
                )
            )
    return findings


def _forbidden_text_findings(path: Path, text: str, validation_contract: dict[str, object]) -> list[ArtifactFinding]:
    findings: list[ArtifactFinding] = []
    for token in _forbidden_string_items(validation_contract):
        if token in text:
            findings.append(
                ArtifactFinding(
                    code="ARTIFACT_FORBIDDEN_TEXT_PRESENT",
                    severity="hard",
                    message="Artifact text contains a forbidden string from validation_contract.",
                    location=str(path),
                    value=token[:200],
                )
            )
    for pattern in _forbidden_regex_items(validation_contract):
        try:
            matched = re.search(pattern, text, flags=re.MULTILINE) is not None
        except re.error as exc:
            findings.append(
                ArtifactFinding(
                    code="ARTIFACT_FORBIDDEN_REGEX_INVALID",
                    severity="hard",
                    message=f"validation_contract forbidden regex is invalid: {exc}",
                    location=str(path),
                    value=pattern[:200],
                )
            )
            continue
        if matched:
            findings.append(
                ArtifactFinding(
                    code="ARTIFACT_FORBIDDEN_REGEX_MATCHED",
                    severity="hard",
                    message="Artifact text matches a forbidden regex from validation_contract.",
                    location=str(path),
                    value=pattern[:200],
                )
            )
    return findings


def _required_string_items(validation_contract: dict[str, object]) -> list[str]:
    return _string_items(
        validation_contract,
        ("required_strings",),
    )


def _required_regex_items(validation_contract: dict[str, object]) -> list[str]:
    return _string_items(
        validation_contract,
        ("required_regex",),
    )


def _forbidden_string_items(validation_contract: dict[str, object]) -> list[str]:
    return _string_items(
        validation_contract,
        ("forbidden_strings",),
    )


def _forbidden_regex_items(validation_contract: dict[str, object]) -> list[str]:
    return _string_items(
        validation_contract,
        ("forbidden_regex",),
    )


def _string_items(validation_contract: dict[str, object], keys: tuple[str, ...]) -> list[str]:
    items: list[str] = []
    for key in keys:
        value = validation_contract.get(key)
        if isinstance(value, str):
            if value.strip():
                items.append(value)
            continue
        if isinstance(value, list | tuple | set):
            items.extend(str(item) for item in value if str(item).strip())
    return list(dict.fromkeys(items))


def _validate_xlsx(
    path: Path,
    validation_contract: dict[str, object] | None = None,
    *,
    workspace_root: Path | None = None,
) -> ArtifactAcceptanceReport:
    try:
        with ZipFile(path) as workbook:
            names = set(workbook.namelist())
    except (BadZipFile, OSError) as exc:
        finding = ArtifactFinding(code="XLSX_INVALID", severity="hard", message=f"Invalid XLSX package: {exc}")
        return _report_with_finding(path, "xlsx", finding)
    missing_parts = _missing_xlsx_parts(names)
    if missing_parts:
        finding = ArtifactFinding(
            code="XLSX_INVALID_PACKAGE",
            severity="hard",
            message="XLSX package is missing required workbook parts.",
            value=",".join(missing_parts),
        )
        return _report_with_finding(path, "xlsx", finding)
    open_finding = _xlsx_open_finding(path)
    if open_finding is not None:
        return _report_with_finding(path, "xlsx", open_finding)
    findings = advisory_artifact_findings([
        *xlsx_contract_findings(path, validation_contract),
        *staged_source_evidence_findings(validation_contract or {}, workspace_root or path.parent),
        *collection_contract_findings(validation_contract or {}, workspace_root or path.parent),
    ])
    return ArtifactAcceptanceReport(
        ok=True,
        artifact_ref=str(path),
        artifact_kind="xlsx",
        findings=findings,
    )


def _validate_pdf(
    path: Path,
    validation_contract: dict[str, object] | None = None,
    *,
    workspace_root: Path | None = None,
) -> ArtifactAcceptanceReport:
    data = path.read_bytes()
    if not data.startswith(b"%PDF-") or b"%%EOF" not in data[-2048:]:
        finding = ArtifactFinding(
            code="PDF_INVALID_SIGNATURE",
            severity="hard",
            message="PDF is missing %PDF header or EOF marker.",
        )
        return _report_with_finding(path, "pdf", finding)
    findings = advisory_artifact_findings([
        *collection_contract_findings(validation_contract or {}, workspace_root or path.parent),
        *document_quality_artifact_findings(path, validation_contract, workspace_root=workspace_root or path.parent),
    ])
    return ArtifactAcceptanceReport(
        ok=True,
        artifact_ref=str(path),
        artifact_kind="pdf",
        findings=findings,
    )


def _missing_xlsx_parts(names: set[str]) -> list[str]:
    missing: list[str] = []
    if "[Content_Types].xml" not in names:
        missing.append("[Content_Types].xml")
    if "xl/workbook.xml" not in names:
        missing.append("xl/workbook.xml")
    if not any(name.startswith("xl/worksheets/") and name.endswith(".xml") for name in names):
        missing.append("xl/worksheets/*.xml")
    return missing


def _xlsx_open_finding(path: Path) -> ArtifactFinding | None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if not workbook.sheetnames:
                return ArtifactFinding("XLSX_NO_VISIBLE_SHEETS", "hard", "XLSX workbook has no visible sheets.")
        finally:
            workbook.close()
    except Exception as exc:
        return ArtifactFinding(
            code="XLSX_INVALID_PACKAGE",
            severity="hard",
            message=f"XLSX cannot be opened by the workbook reader: {exc}",
        )
    return None


def _validate_generic(path: Path) -> ArtifactAcceptanceReport:
    if path.stat().st_size <= 0:
        finding = ArtifactFinding(code="ARTIFACT_EMPTY", severity="hard", message="Artifact is empty.")
        return _report_with_finding(path, kind_for_path(path), finding)
    finding = binary_signature_finding(path)
    if finding is not None:
        return _report_with_finding(path, kind_for_path(path), finding)
    return ArtifactAcceptanceReport(ok=True, artifact_ref=str(path), artifact_kind=kind_for_path(path))


def _report_with_finding(path: Path, kind: str, finding: ArtifactFinding) -> ArtifactAcceptanceReport:
    return ArtifactAcceptanceReport(ok=False, artifact_ref=str(path), artifact_kind=kind, findings=[finding])


def _finding_records(items: list[dict[str, str]]) -> list[ArtifactFinding]:
    return [ArtifactFinding(**item) for item in items]


__all__ = [
    "ArtifactAcceptanceReport",
    "ArtifactAcceptanceRequest",
    "ArtifactFinding",
    "artifact_ref_payload",
    "validate_artifact",
    "validate_html_artifact",
    "validation_workspace_root_for_item",
]
