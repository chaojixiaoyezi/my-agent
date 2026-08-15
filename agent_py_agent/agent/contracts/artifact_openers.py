
from __future__ import annotations

"""产物格式打开器注册表（第 1 层）。

每个打开器只做一件事：把文件 bytes 变成结构化视图，或在打不开时返回一个
结构化的"打不开"finding。打开器**不做内容合同校验**——那是第 2 层通用检查器的事。
未登记的格式没有打开器，验收链路据此落回第 0 层（存在/非空/残桩，格式无关），
绝不"不在表里就拒绝"。

先构造结构化表示，再独立校验；打开器与内容验证正交。
新增一种简单格式 ≈ 在 OPENERS 注册一行；带格式专属深度校验的 ≈ 再加一个薄检查函数。
"""

import csv
import importlib.util
import json
import os
from dataclasses import dataclass, field
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from ..runtime_errors import runtime_error_report
from .artifact_acceptance_models import ArtifactFinding
from .artifact_xlsx_contract import workbook_text, worksheet_tables


@dataclass(frozen=True)
class TabularView:
    """表格类格式（csv/xlsx）的结构化视图。"""

    rows: list[list[str]] = field(default_factory=list)        # csv 解析后的行
    names: set[str] = field(default_factory=set)               # xlsx zip 成员名
    text: str = ""                                             # xlsx 全文本（列名匹配用）
    tables: list[list[list[str]]] = field(default_factory=list)  # xlsx 每个 sheet 的行


@dataclass(frozen=True)
class JsonView:
    """JSON 解析后的顶层值。"""

    value: object


@dataclass(frozen=True)
class GenericView:
    """未登记格式的兜底视图：尽量按"长相"产出可检查的结构。

    text 为能解码出的文本（二进制则为空）；zip_names 为 zip 包成员名（xlsx/docx
    都是 zip，未来未知 zip 格式也能查必含成员）；is_binary 标记是否为二进制。
    通用检查器据此做声明字段校验（min_size / required_sections / required_strings 等），
    不针对任何具体格式，永不随格式增长。
    """

    text: str = ""
    zip_names: set[str] = field(default_factory=set)
    is_binary: bool = False


@dataclass(frozen=True)
class OpenedArtifact:
    """打开结果：opened=True 带 view；否则带 finding（打不开）。

    view 为 None 且 finding 为 None 表示该格式没有结构化视图（如 pdf/docx/txt/md
    的质量校验直接读文件），打开成功但不产视图。
    """

    kind: str
    view: object | None = None
    finding: ArtifactFinding | None = None

    @property
    def opened(self) -> bool:
        return self.finding is None


def open_csv(path: Path) -> OpenedArtifact:
    try:
        rows = list(csv.reader(path.read_text(encoding="utf-8-sig").splitlines()))
    except csv.Error as exc:
        return OpenedArtifact("csv", finding=ArtifactFinding(code="CSV_INVALID", severity="hard", message=f"Invalid CSV: {exc}"))
    return OpenedArtifact("csv", view=TabularView(rows=rows))


def open_xlsx(path: Path) -> OpenedArtifact:
    try:
        with ZipFile(path) as workbook:
            names = set(workbook.namelist())
            text = workbook_text(workbook, names)
            tables = worksheet_tables(workbook, names)
    except (BadZipFile, OSError) as exc:
        return OpenedArtifact("xlsx", finding=ArtifactFinding(code="XLSX_INVALID", severity="hard", message=f"Invalid XLSX package: {exc}"))
    missing_parts = _missing_xlsx_parts(names)
    if missing_parts:
        return OpenedArtifact("xlsx", finding=ArtifactFinding(
            code="XLSX_INVALID_PACKAGE",
            severity="hard",
            message="XLSX package is missing required workbook parts.",
            value=",".join(missing_parts),
        ))
    open_finding = _xlsx_workbook_open_finding(path)
    if open_finding is not None:
        return OpenedArtifact("xlsx", finding=open_finding)
    return OpenedArtifact("xlsx", view=TabularView(names=names, text=text, tables=tables))


def open_pdf(path: Path) -> OpenedArtifact:
    data = path.read_bytes()
    if not data.startswith(b"%PDF-") or b"%%EOF" not in data[-2048:]:
        return OpenedArtifact("pdf", finding=ArtifactFinding(
            code="PDF_INVALID_SIGNATURE",
            severity="hard",
            message="PDF is missing %PDF header or EOF marker.",
        ))
    return OpenedArtifact("pdf")


def open_docx(path: Path) -> OpenedArtifact:
    try:
        with ZipFile(path) as docx:
            names = set(docx.namelist())
    except (BadZipFile, OSError) as exc:
        return OpenedArtifact("docx", finding=ArtifactFinding("DOCX_INVALID", "hard", f"Invalid DOCX package: {exc}"))
    if "word/document.xml" not in names:
        return OpenedArtifact("docx", finding=ArtifactFinding("DOCX_MISSING_DOCUMENT_XML", "hard", "DOCX lacks word/document.xml."))
    return OpenedArtifact("docx")


def open_json(path: Path) -> OpenedArtifact:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        return OpenedArtifact("json", finding=ArtifactFinding(
            code="JSON_INVALID",
            severity="hard",
            message=f"Invalid JSON: {exc.msg}",
            location=str(exc.pos),
        ))
    if not isinstance(value, (dict, list)):
        return OpenedArtifact("json", finding=ArtifactFinding(
            code="JSON_UNEXPECTED_TOP_LEVEL",
            severity="hard",
            message="JSON top-level must be object or array.",
        ))
    return OpenedArtifact("json", view=JsonView(value=value))


def _missing_xlsx_parts(names: set[str]) -> list[str]:
    missing: list[str] = []
    if "[Content_Types].xml" not in names:
        missing.append("[Content_Types].xml")
    if "xl/workbook.xml" not in names:
        missing.append("xl/workbook.xml")
    if not any(name.startswith("xl/worksheets/") and name.endswith(".xml") for name in names):
        missing.append("xl/worksheets/*.xml")
    return missing


def _xlsx_workbook_open_finding(path: Path) -> ArtifactFinding | None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    try:
        return _xlsx_sheet_presence_finding(load_workbook(path, read_only=True, data_only=True))
    except Exception as exc:
        return ArtifactFinding(
            code="XLSX_INVALID_PACKAGE",
            severity="hard",
            message=f"XLSX cannot be opened by the workbook reader: {exc}",
        )


def _xlsx_sheet_presence_finding(workbook) -> ArtifactFinding | None:
    try:
        if not workbook.sheetnames:
            return ArtifactFinding("XLSX_NO_VISIBLE_SHEETS", "hard", "XLSX workbook has no visible sheets.")
        return None
    finally:
        workbook.close()


def open_fallback(path: Path) -> OpenedArtifact:
    """未登记格式的通用兜底打开器：按"长相"产出 GenericView，不拒绝。

    zip 包暴露成员名；能解码的字节当文本；否则标记二进制。打开本身不会失败
    （读不出字节才退回上层第 0 层），所以这里只产视图、不产打不开 finding。
    """
    try:
        data = path.read_bytes()
    except OSError:
        return OpenedArtifact("generic")
    if data[:4] in (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"):
        return OpenedArtifact("generic", view=_zip_member_view(path))
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        return OpenedArtifact("generic", view=GenericView(is_binary=True))
    return OpenedArtifact("generic", view=GenericView(text=text))


def _zip_member_view(path: Path) -> GenericView:
    try:
        with ZipFile(path) as archive:
            return GenericView(zip_names=set(archive.namelist()))
    except (BadZipFile, OSError):
        return GenericView(is_binary=True)


# ---------------------------------------------------------------------------
# 打开器注册表（内置 + 运行时发现的插件）
# ---------------------------------------------------------------------------

# 内置打开器（写死在代码里的已知格式）。
_BUILTIN_OPENERS = {
    "csv": open_csv,
    "xlsx": open_xlsx,
    "pdf": open_pdf,
    "docx": open_docx,
    "json": open_json,
}

# 运行时从打开器目录发现的插件打开器（不改源码、不重启即可扩展）。
_PLUGIN_OPENERS: dict[str, object] = {}

# 插件打开器加载过程中的结构化错误（坏脚本不污染主链路，只记录供观测）。
OPENER_LOAD_ERRORS: list[dict[str, object]] = []

_DISCOVERED = False

# 兼容旧引用：OPENERS 始终指向内置表（已发现插件通过 opener_for 合并查询）。
OPENERS = _BUILTIN_OPENERS


def register_opener(kind: str, opener) -> None:
    """注册一个打开器。插件脚本在 register(register_opener) 钩子里调用本函数。

    只接受 callable；非法登记被忽略并记入 OPENER_LOAD_ERRORS，不抛错。
    内置格式不可被插件覆盖（保证已知格式行为可靠不变）。
    """
    key = str(kind or "").strip().lower()
    if not key or not callable(opener):
        OPENER_LOAD_ERRORS.append({"kind": str(kind), "error": "invalid opener registration (kind empty or not callable)"})
        return
    if key in _BUILTIN_OPENERS:
        OPENER_LOAD_ERRORS.append({"kind": key, "error": "cannot override built-in opener"})
        return
    _PLUGIN_OPENERS[key] = opener


def openers_plugin_dir() -> Path:
    """打开器插件目录：环境变量 MY_AGENT_OPENERS_DIR 优先，否则 ~/.my-agent/openers。"""
    override = os.environ.get("MY_AGENT_OPENERS_DIR", "").strip()
    if override:
        return Path(override).expanduser()
    return Path.home() / ".my-agent" / "openers"


def discover_openers(openers_dir: Path | None = None, *, force: bool = False) -> None:
    """扫描打开器目录，动态加载插件打开器。

    每个 .py 脚本独立隔离加载：导出 `register(register_opener)` 钩子或 `OPENERS`
    dict 即被采纳；任何脚本的导入/执行异常只记进 OPENER_LOAD_ERRORS，绝不让坏脚本
    打断主链路。默认只发现一次（force=True 可重扫，用于热刷新）。
    """
    global _DISCOVERED
    if _DISCOVERED and not force:
        return
    _DISCOVERED = True
    directory = openers_dir or openers_plugin_dir()
    try:
        scripts = sorted(directory.glob("*.py")) if directory.is_dir() else []
    except OSError as exc:
        OPENER_LOAD_ERRORS.append({"path": str(directory), **runtime_error_report(exc, context="artifact_openers.scan")})
        return
    for script in scripts:
        _load_opener_script(script)


def _load_opener_script(script: Path) -> None:
    try:
        spec = importlib.util.spec_from_file_location(f"_my_agent_opener_{script.stem}", script)
        if spec is None or spec.loader is None:
            OPENER_LOAD_ERRORS.append({"path": str(script), "error": "cannot build import spec"})
            return
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
    except Exception as exc:  # 插件脚本任意异常都隔离
        OPENER_LOAD_ERRORS.append({"path": str(script), **runtime_error_report(exc, context="artifact_openers.load")})
        return
    register_hook = getattr(module, "register", None)
    if callable(register_hook):
        try:
            register_hook(register_opener)
        except Exception as exc:
            OPENER_LOAD_ERRORS.append({"path": str(script), **runtime_error_report(exc, context="artifact_openers.register")})
        return
    declared = getattr(module, "OPENERS", None)
    if isinstance(declared, dict):
        for kind, opener in declared.items():
            register_opener(kind, opener)
        return
    OPENER_LOAD_ERRORS.append({"path": str(script), "error": "script exports neither register() nor OPENERS dict"})


def opener_for(kind: str):
    """返回登记的打开器（内置优先，其次运行时发现的插件）。

    未登记返回 None；上层据此走通用兜底（open_fallback）或回到第 0 层。
    """
    key = str(kind or "").strip().lower()
    builtin = _BUILTIN_OPENERS.get(key)
    if builtin is not None:
        return builtin
    discover_openers()
    return _PLUGIN_OPENERS.get(key)


__all__ = [
    "GenericView",
    "JsonView",
    "OpenedArtifact",
    "TabularView",
    "OPENERS",
    "OPENER_LOAD_ERRORS",
    "discover_openers",
    "open_csv",
    "open_docx",
    "open_fallback",
    "open_json",
    "open_pdf",
    "open_xlsx",
    "opener_for",
    "openers_plugin_dir",
    "register_opener",
]
